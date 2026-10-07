"""Read-only Excel view of the BOM, priced at JLCPCB and at a European distributor.
usage: python tools/bom.py [--project kicad/] [--boards 1,5,10] [--out out/bom.xlsx] [--offline] [--refresh]
Needs openpyxl (python -m pip install --user openpyxl).

Outputs (all in out/, git-ignored, regenerate any time):
- bom.xlsx: the priced view.
- bom_jlcpcb.csv: JLCPCB assembly upload (Comment, Designator, Footprint, JLCPCB Part #), fitted parts only.
- bom_order_<N>.csv: Manufacturer, MPN, quantity for the largest board count N, for the TME / Farnell BOM upload.

Source of truth: the KiCad schematic (kicad-cli BOM export, one row per reference) plus docs/external_parts.csv
(modules and other parts bought for the board but not on the schematic). The output is generated, not committed.

Suppliers
- JLCPCB (assembly): basic / extended library, stock, price breaks in USD. No key needed.
- Europe, first one that knows the MPN from the same manufacturer: TME, then Farnell, then Mouser. Each needs a
  free API key in the environment, or its lookups are skipped ("no key"):
    TME_TOKEN, TME_SECRET   (developers.tme.eu, anonymous token + application secret), TME_COUNTRY (default PL)
    FARNELL_API_KEY         (partner.element14.com), FARNELL_STORE (default de.farnell.com)
  JLC_EXTENDED_FEE_USD sets the setup fee per extended line (default 3).
    MOUSER_API_KEY          (mouser.com, Search API)
  RS has no public price API, so it is not included.
Answers are cached in out/bom_cache.json for 7 days. Only real answers and "not found" are cached, so adding a
key or fixing the network takes effect on the next run. --refresh ignores the cache, --offline uses it only.

Quantities per line and board count (need = qty per board x boards):
- JLCPCB: unit price of the price break that covers the need, line = unit x need (JLCPCB adds attrition itself).
- Europe: need rounded up to the supplier minimum and multiple, or a higher price break if that makes the line
  cheaper. The order quantity is shown, line = unit x order qty.

Trust checks (sheet "Check"): the written sheet is read back and its references compared with the KiCad export and
the netlist, every line holds one LCSC number, footprint and manufacturer, and lines without an MPN or a supplier
answer are listed.
"""
import argparse
import base64
import collections
import csv
import hashlib
import hmac
import json
import math
import os
import re
import subprocess
import sys
import tempfile
import time
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sch_check import find_root_sch, kicad_cli  # noqa: E402

SCH = None          # root schematic, set in main()
EXTERNAL = os.path.join(ROOT, "docs", "external_parts.csv")
KICAD_CLI = None    # set in main()
UA = {"User-Agent": "Mozilla/5.0 kicad-ai-template bom tool"}
CACHE_DAYS = 7
EXTENDED_FEE_USD = float(os.environ.get("JLC_EXTENDED_FEE_USD", "3"))   # per extended line per order, check the live quote
TME_COUNTRY = os.environ.get("TME_COUNTRY", "PL")                    # TME prices and stock for this country
FARNELL_STORE = os.environ.get("FARNELL_STORE", "de.farnell.com")     # element14 store id
CATEGORY_ORDER = ["Semiconductors", "Passives", "Connectors", "Protection", "Other", "External modules"]


# ------------------------------------------------------------------ schematic
def kicad(*args):
    r = subprocess.run([KICAD_CLI, *args], capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit(f"kicad-cli failed: {' '.join(args)}\n{r.stderr or r.stdout}")


def read_schematic():
    """One row per reference, all fields we need, DNP included (marked). Also the netlist reference set."""
    tmp = tempfile.mkdtemp()
    bom_csv = os.path.join(tmp, "bom.csv")
    fields = "Reference,Value,Footprint,Manufacturer,MPN,LCSC Part,Description,${DNP}"
    kicad("sch", "export", "bom", "--fields", fields, "--labels", "Ref,Value,Footprint,Manufacturer,MPN,LCSC,Description,DNP",
          "--group-by", "", "--ref-range-delimiter", "", "-o", bom_csv, SCH)
    with open(bom_csv, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    net = os.path.join(tmp, "net.xml")
    kicad("sch", "export", "netlist", "--format", "kicadxml", "-o", net, SCH)
    comps = list(ET.parse(net).getroot().iter("comp"))
    net_refs = {c.get("ref") for c in comps}
    # parts marked "exclude from BOM" (e.g. bare solder pads) are left out of the BOM export by kicad-cli
    excluded = {c.get("ref") for c in comps if any(p.get("name") == "exclude_from_bom" for p in c.iter("property"))}
    # an excluded part that has a part number is probably excluded by mistake
    excluded_with_pn = sorted(c.get("ref") for c in comps if c.get("ref") in excluded and any(
        p.get("name") in ("MPN", "LCSC Part") and (p.get("value") or "").strip() for p in c.iter("property")))
    return rows, net_refs, excluded, excluded_with_pn


CATEGORY = {"R": "Passives", "C": "Passives", "L": "Passives", "FB": "Passives",
            "D": "Semiconductors", "Q": "Semiconductors", "U": "Semiconductors",
            "J": "Connectors", "SW": "Connectors", "F": "Protection"}


def category(ref):
    return CATEGORY.get(re.match(r"[A-Za-z]+", ref).group(0), "Other")


def ref_key(ref):
    m = re.match(r"([A-Za-z]+)(\d+)", ref)
    return (m.group(1), int(m.group(2))) if m else (ref, 0)


def group(rows):
    """Lines = same MPN (or same value + footprint when the MPN is empty), split by DNP. Sorted by category."""
    lines = collections.OrderedDict()
    for r in sorted(rows, key=lambda r: ref_key(r["Ref"])):
        dnp = r["DNP"].strip() not in ("", "0", "false", "False")
        key = (r["MPN"].strip() or f'{r["Value"]}|{r["Footprint"]}', dnp)
        lines.setdefault(key, []).append(r)
    out = []
    for rs in lines.values():
        r0 = rs[0]
        out.append({"category": category(r0["Ref"]), "refs": [r["Ref"] for r in rs], "qty": len(rs),
                    "value": ", ".join(sorted({r["Value"] for r in rs})), "footprint": r0["Footprint"].split(":")[-1],
                    "manufacturer": r0["Manufacturer"], "mpn": r0["MPN"].strip(), "lcsc": r0["LCSC"].strip(),
                    "description": r0["Description"], "dnp": r0["DNP"].strip() not in ("", "0", "false", "False"),
                    "external": False, "est_eur": None,
                    "mixed": {k: sorted({r[k] for r in rs}) for k in ("LCSC", "Footprint", "Manufacturer")
                              if len({r[k] for r in rs}) > 1}})
    out.sort(key=lambda ln: (CATEGORY_ORDER.index(ln["category"]), ref_key(ln["refs"][0])))
    return out


def read_external():
    out = []
    if not os.path.exists(EXTERNAL):
        return out
    with open(EXTERNAL, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            out.append({"category": r.get("Category") if r.get("Category") in CATEGORY_ORDER else "External modules", "refs": [], "qty": int(r["Qty_per_board"]),
                        "value": r["Item"], "footprint": "", "manufacturer": r["Manufacturer"], "mpn": r["MPN"].strip(),
                        "lcsc": "", "description": r["Notes"], "dnp": False, "external": True, "mixed": {},
                        "est_eur": float(r["Est_price_EUR"]) if r["Est_price_EUR"] else None})
    return out


# ------------------------------------------------------------------ price helpers
def _pl(n):
    return "" if n == 1 else "s"


def unit_at(breaks, q):
    """Unit price of the price break that covers quantity q (breaks: [(qty, unit)] ascending)."""
    u = breaks[0][1]
    for bq, p in breaks:
        if q >= bq:
            u = p
    return u


def eu_order(need, breaks, moq=1, mult=1):
    """Cheapest (order qty, unit, line) that covers need: minimum, multiple, or a higher break if cheaper."""
    moq, mult = max(1, int(moq or 1)), max(1, int(mult or 1))

    def rounded(q):
        return int(math.ceil(max(q, moq) / mult) * mult)

    cands = {rounded(need)} | {rounded(bq) for bq, _ in breaks if bq > need}
    return min(((q, unit_at(breaks, q), q * unit_at(breaks, q)) for q in cands), key=lambda t: (t[2], t[0]))


def jlc_line(ln, n):
    j = ln.get("jlc") or {}
    if ln["dnp"] or not j.get("breaks"):
        return None
    need = ln["qty"] * n
    u = unit_at(j["breaks"], need)
    return need, u, need * u


def eu_line(ln, n):
    e = ln.get("eu") or {}
    if ln["dnp"] or not e.get("breaks"):
        return None
    return eu_order(ln["qty"] * n, e["breaks"], e.get("moq", 1), e.get("mult", 1))


def http(url, data=None, headers=None):
    req = urllib.request.Request(url, data=data, headers={**UA, **(headers or {})})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))


class Cache:
    """Real answers and "not found" are kept for CACHE_DAYS. Errors (no key, timeouts, API errors) are not kept."""

    def __init__(self, path, offline, refresh):
        self.path, self.offline = path, offline
        try:
            self.d = {} if refresh else json.load(open(path, encoding="utf-8"))
        except (OSError, ValueError):
            self.d = {}
        self.oldest = None

    def get(self, key, fn):
        hit = self.d.get(key)
        if not (isinstance(hit, dict) and "t" in hit and "v" in hit):     # older cache format: refetch
            hit = None
        if hit and (self.offline or time.time() - hit["t"] < CACHE_DAYS * 86400):
            self.oldest = min(self.oldest or hit["t"], hit["t"])
            return dict(hit["v"])
        if self.offline:
            return {"error": "offline, not cached"}
        try:
            v = fn()
        except Exception as e:                          # one failed lookup must not stop the sheet
            v = {"error": f"{type(e).__name__}: {e}"[:200]}
        if "error" not in v or v["error"] == "not found":
            self.d[key] = {"t": time.time(), "v": v}
            time.sleep(0.3)
        return dict(v)

    def save(self):
        os.makedirs(os.path.dirname(self.path), exist_ok=True)
        json.dump(self.d, open(self.path, "w", encoding="utf-8"), indent=1)


def num(s):
    """'€0,123' / '0.12' / 1.2 -> float."""
    if isinstance(s, (int, float)):
        return float(s)
    s = re.sub(r"[^\d,.\-]", "", str(s))
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    return float(s) if s else None


def same_maker(want, got):
    """Loose manufacturer match: the first word of one name appears in the other (YAGEO vs Yageo Corporation)."""
    w, g = (want or "").lower(), (got or "").lower()
    if not w or not g:
        return True
    return w.split()[0] in g or g.split()[0] in w


# ------------------------------------------------------------------ JLCPCB
def jlc(lcsc):
    d = http("https://jlcpcb.com/api/overseas-pcb-order/v1/shoppingCart/smtGood/selectSmtComponentList",
             json.dumps({"keyword": lcsc, "currentPage": 1, "pageSize": 5}).encode(), {"Content-Type": "application/json"})
    if d.get("code") != 200:
        return {"error": f'JLCPCB API: {d.get("message") or d.get("code")}'}
    for c in ((d.get("data") or {}).get("componentPageInfo") or {}).get("list") or []:
        if c.get("componentCode") == lcsc:
            return {"library": ("extended (preferred)" if c.get("componentLibraryType") == "expand" and c.get("preferredComponentFlag")
                                else {"base": "basic", "expand": "extended"}.get(c.get("componentLibraryType"), c.get("componentLibraryType"))),
                    "stock": c.get("stockCount"), "currency": "USD",
                    "breaks": [(p["startNumber"], p["productPrice"]) for p in c.get("componentPrices") or []],
                    "url": f"https://jlcpcb.com/partdetail/{lcsc}"}
    return {"error": "not found"}


# ------------------------------------------------------------------ Europe
def tme(mpn, maker):
    tok, sec = os.environ.get("TME_TOKEN"), os.environ.get("TME_SECRET")
    if not (tok and sec):
        return {"error": "no key"}

    def call(action, params):
        url = f"https://api.tme.eu/{action}.json"
        params = {**params, "Token": tok, "Country": TME_COUNTRY, "Language": "EN", "Currency": "EUR"}
        q = lambda s, *a: urllib.parse.quote(s, safe="")          # encode "/" too, the signature needs it
        enc = urllib.parse.urlencode(sorted(params.items()), quote_via=q)
        base = "POST&" + q(url) + "&" + q(enc)
        sig = base64.b64encode(hmac.new(sec.encode(), base.encode(), hashlib.sha1).digest()).decode()
        d = http(url, (enc + "&ApiSignature=" + q(sig)).encode(), {"Content-Type": "application/x-www-form-urlencoded"})
        if d.get("Status") != "OK":
            raise RuntimeError(f'TME API {d.get("Status")}: {d.get("Error", "")}')
        return d

    prods = (call("Products/Search", {"SearchPlain": mpn}).get("Data") or {}).get("ProductList") or []
    prod = next((p for p in prods if p.get("OriginalSymbol", "").upper() == mpn.upper()
                 and same_maker(maker, p.get("Producer"))), None)
    if not prod:
        return {"error": "not found"}
    p = ((call("Products/GetPricesAndStocks", {"SymbolList[0]": prod["Symbol"]}).get("Data") or {}).get("ProductList") or [{}])[0]
    return {"supplier": "TME", "part": prod["Symbol"], "stock": p.get("Amount"), "moq": prod.get("MinAmount", 1),
            "mult": prod.get("Multiples", 1), "currency": "EUR", "lead": "",
            "breaks": [(b["Amount"], b["PriceValue"]) for b in p.get("PriceList") or []],
            "url": "https://www.tme.eu/en/details/" + urllib.parse.quote(prod["Symbol"].lower())}


def farnell(mpn, maker):
    key = os.environ.get("FARNELL_API_KEY")
    if not key:
        return {"error": "no key"}
    q = urllib.parse.urlencode({"term": f"manuPartNum:{mpn}", "storeInfo.id": FARNELL_STORE,
                                "resultsSettings.offset": 0, "resultsSettings.numberOfResults": 10,
                                "resultsSettings.responseGroup": "large,inventory", "callInfo.responseDataFormat": "json",
                                "callInfo.apiKey": key})
    d = http("https://api.element14.com/catalog/products?" + q)
    if "Fault" in d:
        raise RuntimeError(f'Farnell API: {d["Fault"]}')
    prods = (d.get("manufacturerPartNumberSearchReturn") or {}).get("products") or []
    p = next((p for p in prods if same_maker(maker, p.get("vendorName") or p.get("brandName"))), None)
    if not p:
        return {"error": "not found"}
    st = p.get("stock") or {}
    return {"supplier": "Farnell", "part": p.get("sku"), "stock": st.get("level"),
            "moq": p.get("translatedMinimumOrderQuality", 1), "mult": p.get("orderMultiple") or 1, "currency": "EUR",
            "lead": f'{st.get("leastLeadTime")} days' if st.get("leastLeadTime") else "",
            "breaks": [(b["from"], b["cost"]) for b in p.get("prices") or []],
            "url": f'https://{FARNELL_STORE}/{p.get("sku")}'}


def mouser(mpn, maker):
    key = os.environ.get("MOUSER_API_KEY")
    if not key:
        return {"error": "no key"}
    d = http(f"https://api.mouser.com/api/v1/search/partnumber?apiKey={key}",
             json.dumps({"SearchByPartRequest": {"mouserPartNumber": mpn, "partSearchOptions": "Exact"}}).encode(),
             {"Content-Type": "application/json"})
    if d.get("Errors"):
        raise RuntimeError(f'Mouser API: {d["Errors"]}')
    parts = ((d.get("SearchResults") or {}).get("Parts")) or []
    p = next((p for p in parts if same_maker(maker, p.get("Manufacturer"))), None)
    if not p:
        return {"error": "not found"}
    pb = p.get("PriceBreaks") or []
    stock = num(p.get("AvailabilityInStock") or (p.get("Availability") or "0").split(" ")[0]) or 0
    return {"supplier": "Mouser", "part": p.get("MouserPartNumber"), "stock": int(stock), "moq": int(p.get("Min") or 1),
            "mult": int(p.get("Mult") or 1), "currency": pb[0].get("Currency", "") if pb else "",
            "lead": p.get("LeadTime") or "", "breaks": [(b["Quantity"], num(b["Price"])) for b in pb],
            "url": p.get("ProductDetailUrl")}


EU = [("TME", tme), ("Farnell", farnell), ("Mouser", mouser)]


def eu_lookup(cache, mpn, maker):
    notes = []
    for name, fn in EU:
        r = cache.get(f"{name}:{mpn}", lambda fn=fn: fn(mpn, maker))
        if "error" not in r and r.get("breaks"):
            return r
        notes.append(f'{name}: {r.get("error", "no price")}')
    return {"error": ", ".join(notes)}


# ------------------------------------------------------------------ Excel
def write_xlsx(path, lines, boards, generated, cache_note):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active; ws.title = "BOM"
    bold = Font(bold=True)
    ws["A1"] = f"{os.path.splitext(os.path.basename(SCH))[0]} BOM (read-only view, generated)"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Generated {generated} from the KiCad schematic + docs/external_parts.csv. {cache_note} "
                "JLCPCB prices in USD. European prices in the currency of the EU currency column.")
    ws["A3"] = ("One row = one BOM line: one unique part (one MPN), however many times it is used. "
                f"Qty/board = how many of that part one board needs. JLCPCB charges about ${EXTENDED_FEE_USD:g} per extended line per order "
                "(preferred extended parts: none on Economic PCBA, check the quote).")
    head = ["Category", "Refs", "Qty/board", "Value", "Footprint", "Manufacturer", "MPN", "Description", "DNP",
            "JLC LCSC#", "JLC library", "JLC stock"]
    for n in boards:
        head += [f"JLC unit [USD] ({n} board{_pl(n)})", f"JLC line [USD] ({n})"]
    head += ["EU supplier", "EU part#", "EU stock", "EU lead time", "EU min / mult", "EU currency"]
    for n in boards:
        head += [f"EU order qty ({n} board{_pl(n)})", f"EU unit ({n})", f"EU line ({n})"]
    head += ["External: est. EUR per board", "Notes"]
    R0 = 4
    for c, h in enumerate(head, 1):
        cell = ws.cell(R0, c, h); cell.font = bold; cell.alignment = Alignment(wrap_text=True, vertical="top")
    col = {h: i + 1 for i, h in enumerate(head)}
    L = lambda h: get_column_letter(col[h])
    red, orange, yellow, grey = (PatternFill("solid", fgColor=x) for x in ("F8B4B4", "FCD9A8", "FFF3B0", "E6E6E6"))
    r = R0
    for ln in lines:
        r += 1
        ln["_row"] = r
        v = [ln["category"], ",".join(ln["refs"]), ln["qty"], ln["value"], ln["footprint"], ln["manufacturer"],
             ln["mpn"], ln["description"], "DNP" if ln["dnp"] else ""]
        for c, x in enumerate(v, 1):
            ws.cell(r, c, x)
        j, e = ln.get("jlc") or {}, ln.get("eu") or {}
        if ln["lcsc"]:
            ws.cell(r, col["JLC LCSC#"], ln["lcsc"]).hyperlink = j.get("url")
        ws.cell(r, col["JLC library"], j.get("library") or j.get("error", ""))
        ws.cell(r, col["JLC stock"], j.get("stock"))
        ws.cell(r, col["EU supplier"], e.get("supplier") or e.get("error", ""))
        if e.get("part"):
            ws.cell(r, col["EU part#"], e["part"]).hyperlink = e.get("url")
            ws.cell(r, col["EU stock"], e.get("stock"))
            ws.cell(r, col["EU lead time"], e.get("lead", ""))
            ws.cell(r, col["EU min / mult"], f'{e.get("moq", 1)} / {e.get("mult", 1)}')
            ws.cell(r, col["EU currency"], e.get("currency", ""))
        notes = []
        for n in boards:
            jl = jlc_line(ln, n)
            if jl:
                ws.cell(r, col[f"JLC unit [USD] ({n} board{_pl(n)})"], jl[1])
                ws.cell(r, col[f"JLC line [USD] ({n})"], f'={L(f"JLC unit [USD] ({n} board{_pl(n)})")}{r}*{jl[0]}')
                if j.get("stock") is not None and j["stock"] < jl[0]:
                    ws.cell(r, col["JLC stock"]).fill = red
            el = eu_line(ln, n)
            if el:
                ws.cell(r, col[f"EU order qty ({n} board{_pl(n)})"], el[0])
                ws.cell(r, col[f"EU unit ({n})"], el[1])
                ws.cell(r, col[f"EU line ({n})"], f'={L(f"EU unit ({n})")}{r}*{L(f"EU order qty ({n} board{_pl(n)})")}{r}')
                if e.get("stock") is not None and e["stock"] < el[0]:
                    ws.cell(r, col["EU stock"]).fill = red
        if ln["est_eur"] is not None:
            ws.cell(r, col["External: est. EUR per board"], ln["est_eur"] * ln["qty"])
        if ln["dnp"]:
            notes.append("not fitted, not priced")
            for c in range(1, len(head) + 1):
                ws.cell(r, c).fill = grey
        if ln["mixed"]:
            notes.append("mixed fields: " + ", ".join(f"{k} {v}" for k, v in ln["mixed"].items()))
        ws.cell(r, col["Notes"], ", ".join(notes))
    last = r
    # biggest cost offenders: 1st / 2nd / 3rd most expensive line per category, at the largest board count
    n = max(boards)
    for key, fn in ((f"JLC line [USD] ({n})", jlc_line), (f"EU line ({n})", eu_line)):
        bycat = collections.defaultdict(list)
        for ln in lines:
            x = fn(ln, n)
            if x:
                bycat[ln["category"]].append((x[2], ln["_row"]))
        for items in bycat.values():
            for rank, (_, row) in enumerate(sorted(items, reverse=True)[:3]):
                ws.cell(row, col[key]).fill = (red, orange, yellow)[rank]
    # totals per category and grand total (formulas, so own calculations can build on them)
    r = last + 2
    ws.cell(r, 1, "Totals").font = bold
    A = f"$A${R0 + 1}:$A${last}"
    for cat in [c for c in CATEGORY_ORDER if any(ln["category"] == c for ln in lines)] + ["All"]:
        r += 1
        ws.cell(r, 1, cat).font = bold
        for h in [f"JLC line [USD] ({n})" for n in boards] + [f"EU line ({n})" for n in boards] + ["External: est. EUR per board"]:
            rng = f"{L(h)}{R0 + 1}:{L(h)}{last}"
            ws.cell(r, col[h], f"=SUM({rng})" if cat == "All" else f'=SUMIF({A},"{cat}",{rng})').font = bold
    r += 1
    ws.cell(r, 1, f"JLCPCB extended setup fees, ${EXTENDED_FEE_USD:g} per extended line per order (preferred ones included)").font = bold
    for n in boards:
        ws.cell(r, col[f"JLC line [USD] ({n})"],
                f'=COUNTIFS({L("JLC library")}{R0 + 1}:{L("JLC library")}{last},"extended*",{L("DNP")}{R0 + 1}:{L("DNP")}{last},"")*{EXTENDED_FEE_USD:g}').font = bold
    widths = {"Refs": 28, "Value": 18, "Description": 40, "Notes": 30, "MPN": 22, "Footprint": 22, "EU supplier": 30}
    for h, i in col.items():
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 12)
    ws.freeze_panes = ws.cell(R0 + 1, 3)
    ws.auto_filter.ref = f"A{R0}:{get_column_letter(len(head))}{last}"
    ws.row_dimensions[R0].height = 45

    lg = wb.create_sheet("Legend")
    for i, t in enumerate([
            "Rows are sorted by category: semiconductors, passives, connectors, protection, external modules.",
            "Red / orange / yellow line total: the 1st / 2nd / 3rd most expensive line of its category, at the largest board count.",
            "Red stock cell: the supplier has fewer than the quantity needed for that board count.",
            "Grey row: DNP, not fitted, not priced.",
            "JLC line = unit price of the break that covers the need x need. JLCPCB adds attrition itself.",
            "EU columns: TME, then Farnell, then Mouser, the first with the MPN from the same manufacturer. 'no key' = API key missing.",
            "EU order qty: need rounded up to the supplier minimum and multiple, or more if a higher price break is cheaper.",
            "External parts: rough estimates from docs/external_parts.csv, per board.",
            f"Extended JLCPCB parts add a setup fee (about ${EXTENDED_FEE_USD:g}) each per order: own row under the totals."], 1):
        lg.cell(i, 1, t)
    lg.column_dimensions["A"].width = 140
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return R0


def add_checks(path, checks):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    wb = load_workbook(path)
    ck = wb.create_sheet("Check", 1)
    ck["A1"] = "Trust checks"; ck["A1"].font = Font(bold=True, size=14)
    for i, (name, ok, detail) in enumerate(checks, 3):
        ck.cell(i, 1, name)
        c = ck.cell(i, 2, "OK" if ok else "CHECK")
        c.fill = PatternFill("solid", fgColor="C6EFCE" if ok else "F8B4B4")
        ck.cell(i, 3, detail)
    ck.column_dimensions["A"].width = 46; ck.column_dimensions["C"].width = 120
    wb.save(path)


def read_back(path, R0):
    """References and quantities as they stand in the written sheet."""
    from openpyxl import load_workbook
    ws = load_workbook(path)["BOM"]
    refs, qty_ok, ext = [], True, 0
    for r in range(R0 + 1, ws.max_row + 1):
        cat = ws.cell(r, 1).value
        if cat is None:                                  # blank row before the totals
            break
        cell = ws.cell(r, 2).value or ""
        rr = [x for x in cell.split(",") if x]
        refs += rr
        if cat == "External modules":
            ext += 1
        elif len(rr) != ws.cell(r, 3).value:
            qty_ok = False
    return refs, qty_ok, ext


# ------------------------------------------------------------------ main
def write_csvs(out_dir, lines, n):
    """JLCPCB assembly BOM and a distributor order list, both from the same lines as the sheet."""
    fitted = [ln for ln in lines if not ln["dnp"] and not ln["external"]]
    p1 = os.path.join(out_dir, "bom_jlcpcb.csv")
    with open(p1, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Comment", "Designator", "Footprint", "JLCPCB Part #"])
        for ln in fitted:
            w.writerow([ln["value"], ",".join(ln["refs"]), ln["footprint"], ln["lcsc"]])
    p2 = os.path.join(out_dir, f"bom_order_{n}.csv")
    with open(p2, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Manufacturer", "MPN", "Qty", "Refs", "Value"])
        for ln in fitted:
            w.writerow([ln["manufacturer"], ln["mpn"], ln["qty"] * n, ",".join(ln["refs"]), ln["value"]])
    return p1, p2


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--boards", default="1,5,10")
    ap.add_argument("--out", default=os.path.join(ROOT, "out", "bom.xlsx"))
    ap.add_argument("--offline", action="store_true", help="use the cache only, no network")
    ap.add_argument("--refresh", action="store_true", help="ignore the cache")
    ap.add_argument("--project", help="project folder or .kicad_pro (default: the only project in the repo)")
    a = ap.parse_args()
    global SCH, KICAD_CLI
    SCH, KICAD_CLI = find_root_sch(a.project), kicad_cli()
    boards = [int(x) for x in a.boards.split(",")]

    rows, net_refs, excluded, excluded_with_pn = read_schematic()
    lines = group(rows)
    ext = read_external()
    cache = Cache(os.path.join(ROOT, "out", "bom_cache.json"), a.offline, a.refresh)
    for ln in lines + ext:
        if ln["lcsc"]:
            ln["jlc"] = cache.get(f"JLC:{ln['lcsc']}", lambda ln=ln: jlc(ln["lcsc"]))
        if ln["mpn"]:
            ln["eu"] = eu_lookup(cache, ln["mpn"], ln["manufacturer"])
    cache.save()
    cache_note = ("Supplier data fetched now." if cache.oldest is None else
                  f"Oldest supplier data from {time.strftime('%Y-%m-%d', time.localtime(cache.oldest))} (cache).")

    generated = time.strftime("%Y-%m-%d %H:%M")
    R0 = write_xlsx(a.out, lines + ext, boards, generated, cache_note)

    refs = [r["Ref"] for r in rows]
    sheet_refs, qty_ok, n_ext = read_back(a.out, R0)
    dup = sorted({x for x in sheet_refs if sheet_refs.count(x) > 1})
    fitted = [ln for ln in lines if not ln["dnp"]]
    no_eu = [ln for ln in fitted if not (ln.get("eu") or {}).get("breaks")]
    checks = [
        ("Sheet references = KiCad BOM export", sorted(sheet_refs) == sorted(refs) and not dup,
         f"{len(refs)} in the export, {len(sheet_refs)} in the written sheet. Missing: "
         f"{', '.join(sorted(set(refs) - set(sheet_refs))) or 'none'}. Extra: {', '.join(sorted(set(sheet_refs) - set(refs))) or 'none'}. "
         f"Duplicates: {', '.join(dup) or 'none'}"),
        ("KiCad BOM export = netlist", set(refs) == net_refs - excluded and not set(refs) & excluded and not excluded_with_pn,
         f"only in the BOM export: {', '.join(sorted(set(refs) - net_refs)) or 'none'}. "
         f"only in the netlist: {', '.join(sorted(net_refs - excluded - set(refs))) or 'none'}. "
         f"Excluded from the BOM in the schematic: {', '.join(sorted(excluded)) or 'none'}. "
         f"Exported although excluded: {', '.join(sorted(set(refs) & excluded)) or 'none'}. "
         f"Excluded but with an MPN or LCSC number: {', '.join(excluded_with_pn) or 'none'}"),
        ("Qty per board = references on the line (sheet)", qty_ok, f"{len(lines)} schematic lines"),
        ("One LCSC#, footprint and manufacturer per line", not any(ln["mixed"] for ln in lines),
         ", ".join(f'{",".join(ln["refs"])}: {ln["mixed"]}' for ln in lines if ln["mixed"]) or "none"),
        ("Fitted parts per board", True, f'{sum(ln["qty"] for ln in fitted)} fitted, '
                                          f'{sum(ln["qty"] for ln in lines if ln["dnp"])} DNP '
                                          f'({", ".join(x for ln in lines if ln["dnp"] for x in ln["refs"]) or "none"})'),
        ("External parts in the sheet", n_ext == len(ext), f"{n_ext} in the sheet, {len(ext)} in docs/external_parts.csv"),
        ("Lines without MPN", not [ln for ln in lines if not ln["mpn"]],
         ", ".join(",".join(ln["refs"]) for ln in lines if not ln["mpn"]) or "none"),
        ("Lines without LCSC number (JLCPCB)", not [ln for ln in lines if not ln["lcsc"]],
         ", ".join(",".join(ln["refs"]) for ln in lines if not ln["lcsc"]) or "none"),
        ("Fitted lines JLCPCB could not price", not [ln for ln in fitted if not (ln.get("jlc") or {}).get("breaks")],
         ", ".join(f'{ln["mpn"]} ({(ln.get("jlc") or {}).get("error", "")})' for ln in fitted
                   if not (ln.get("jlc") or {}).get("breaks")) or "none"),
        ("Fitted lines no European supplier could price", not no_eu,
         f'{len(no_eu)} lines. First reason: {(no_eu[0].get("eu") or {}).get("error", "") if no_eu else "none"}'),
    ]
    add_checks(a.out, checks)
    csvs = write_csvs(os.path.dirname(a.out), lines, max(boards))
    print(f"{a.out}: {len(lines)} schematic lines ({len(refs)} references) + {len(ext)} external. {cache_note}")
    for name, ok, detail in checks:
        print(f"  {'OK   ' if ok else 'CHECK'} {name}: {detail[:110]}")
    print("  also written: " + ", ".join(csvs))


if __name__ == "__main__":
    main()
